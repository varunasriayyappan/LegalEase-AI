import html
import io
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from fpdf import FPDF

from config import FOOTER_TEXT, LOGO_PATH, ensure_logo

_REPLACEMENTS = {
    "\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"',
    "\u2013": "-", "\u2014": "-", "\u2022": "-", "\u2026": "...",
    "\u00a0": " ",
}

_HEADING_RE = re.compile(r"^\d+\.\s+[A-Za-z][^.]{0,70}:?$")


def sanitize_text(text: str) -> str:
    """Replace typographic characters and strip Markdown symbols."""
    for old, new in _REPLACEMENTS.items():
        text = text.replace(old, new)
    text = text.replace("**", "").replace("__", "")
    text = re.sub(r"^\s{0,3}#{1,6}\s*", "", text, flags=re.MULTILINE)
    text = re.sub(r"^\s*[*]\s+", "- ", text, flags=re.MULTILINE)
    return text.strip()


def split_terms(terms: str) -> list[str]:
    return [t.strip() for t in (terms or "").split(";") if t.strip()]


def is_heading(line: str) -> bool:
    s = line.strip()
    if not s or len(s) > 80:
        return False
    if _HEADING_RE.match(s):
        return True
    if s.isupper() and len(s) > 3:
        return True
    return s.endswith(":") and len(s) < 60


# ---------------------------------------------------------------- HTML
def format_html_preview(text: str) -> str:
    parts = []
    for line in text.splitlines():
        if not line.strip():
            continue
        safe = html.escape(line.strip())
        if is_heading(line):
            parts.append(
                f"<h4 style='margin:18px 0 6px 0;color:#93c5fd;'>{safe}</h4>"
            )
        else:
            parts.append(f"<p style='margin:6px 0;line-height:1.6;'>{safe}</p>")
    return "".join(parts)


# ---------------------------------------------------------------- DOCX
def format_docx(text: str, doc_type: str, terms: str = "") -> bytes:
    ensure_logo()
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    # Logo
    if LOGO_PATH.exists():
        doc.add_picture(str(LOGO_PATH), width=Inches(2.2))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(doc_type.title())
    run.bold = True
    run.font.size = Pt(16)

    # Body
    for line in text.splitlines():
        if not line.strip():
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(line.strip())
        if is_heading(line):
            run.bold = True
            p.paragraph_format.space_before = Pt(10)

    # Terms table
    items = split_terms(terms)
    if items:
        doc.add_paragraph()
        h = doc.add_paragraph()
        h.add_run("Summary of Key Terms").bold = True
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "No."
        table.rows[0].cells[1].text = "Term / Condition"
        for cell in table.rows[0].cells:
            for r in cell.paragraphs[0].runs:
                r.bold = True
        for i, item in enumerate(items, 1):
            row = table.add_row().cells
            row[0].text = str(i)
            row[1].text = item

    # Footer
    footer_p = doc.sections[0].footer.paragraphs[0]
    footer_p.text = FOOTER_TEXT
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- PDF
def _latin1(s: str) -> str:
    """Built-in PDF fonts only support Latin-1."""
    return s.encode("latin-1", "replace").decode("latin-1")


class _LegalPDF(FPDF):
    def header(self):
        if LOGO_PATH.exists():
            self.image(str(LOGO_PATH), x=(self.w - 40) / 2, y=8, w=40)
        self.set_y(26)

    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica", "I", 8)
        self.cell(0, 10, _latin1(FOOTER_TEXT), align="C")


def format_pdf(text: str, doc_type: str, terms: str = "") -> bytes:
    ensure_logo()
    pdf = _LegalPDF()
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.set_margins(20, 20, 20)
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 15)
    pdf.cell(0, 10, _latin1(doc_type.title()), align="C",
             new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)

    for line in text.splitlines():
        s = _latin1(line.strip())
        if not s:
            pdf.ln(2)
            continue
        if is_heading(line):
            pdf.ln(2)
            pdf.set_font("Helvetica", "B", 11)
        else:
            pdf.set_font("Helvetica", "", 10.5)
        pdf.multi_cell(0, 6, s, new_x="LMARGIN", new_y="NEXT")

    items = split_terms(terms)
    if items:
        pdf.ln(4)
        pdf.set_font("Helvetica", "B", 11)
        pdf.cell(0, 8, "Summary of Key Terms", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 10.5)
        for item in items:
            pdf.multi_cell(0, 6, _latin1("- " + item),
                           new_x="LMARGIN", new_y="NEXT")

    return bytes(pdf.output())